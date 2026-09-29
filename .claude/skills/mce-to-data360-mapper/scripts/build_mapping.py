#!/usr/bin/env python3
"""Build the MCE -> Data 360 mapping workbook and SFMC inventory CSVs.

Usage: build_mapping.py shortlist.json fields.json out_dir [--prefix CLIENT] [--rules mapping_rules.json]
Outputs: <prefix>_SFMC_DE_summary.csv, <prefix>_SFMC_DE_field_inventory.csv,
         <prefix>_MCE_to_Data360_Mapping.xlsx (5 tabs). Prints tier counts JSON to stdout.
"""
import json,re,sys,os,csv
from collections import defaultdict,Counter
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
a=sys.argv; short,fpath,out=a[1],a[2],a[3]
prefix=a[a.index('--prefix')+1] if '--prefix' in a else 'Client'
rules=json.load(open(a[a.index('--rules')+1] if '--rules' in a else os.path.join(os.path.dirname(__file__),'mapping_rules.json')))
os.makedirs(out,exist_ok=True)
des={d['key']:d for d in json.load(open(short))}; fields=json.load(open(fpath))
byde=defaultdict(list)
for r in fields: byde[r['key']].append(r)

def tier(d):
    for t in rules['tiers']:
        if 'name' in t and not re.search(t['name'],d['name']): continue
        if 'path' in t and not re.search(t['path'],d['path']): continue
        if 'name' not in t and 'path' not in t: continue
        if t.get('skip_if') and re.search(t['skip_if'],d['name']): return ('X','Skip','Operational config / log for '+t['tier'],'')
        target=rules['personalization_targets'].get(d['name'],'') if t['target']=='see personalization_targets' else t['target']
        return (t['tier'],t['approach'],t['rationale'],target)
    return ('X','Skip','No rule matched. Review manually.','')
for d in des.values(): d['tier'],d['approach'],d['rationale'],d['target']=tier(d)

def map_field(d,r):
    for pat,dmo,f,note in rules['fields']:
        if re.match(pat,r['field'],re.I): return dmo,f,note
    base=d['target'].split(' (')[0].split(' +')[0]
    if d['tier'] in ('H1','H2','F','G'): return base,re.sub(r'[^A-Za-z0-9]+','_',r['field']).strip('_')+'__c','Custom attribute on custom DMO'
    if d['tier'] in ('B','C','D','E'): return base,'(custom attribute)','No standard field. Add custom attribute or drop.'
    return '','','Not mapped'

def layer(p):
    t=p.split(' / '); t=t[1] if len(t)>1 else t[0]
    return {'_Master':'Master data','05 System':'System and sync','03 Journeys':'Journey entry and log','04 Reporting':'Reporting','02 Transactional':'Transactional','Data Extensions':'Platform and personalization'}.get(t,t)

# CSVs
with open(f'{out}/{prefix}_SFMC_DE_field_inventory.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['de_name','customer_key','folder_path','layer','is_sendable','ordinal','field_name','field_type','max_length','is_primary_key','is_required','last_modified'])
    for k,fl in byde.items():
        d=des[k]
        for r in sorted(fl,key=lambda x:int(x['ord'] or 0)): w.writerow([d['name'],k,d['path'],layer(d['path']),'Y' if d['send']=='true' else '',r['ord'],r['field'],r['ftype'],r['maxlen'],'Y' if r['pk']=='true' else '','Y' if r['req']=='true' else '',d['mod']])
with open(f'{out}/{prefix}_SFMC_DE_summary.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['de_name','customer_key','folder_path','layer','field_count','primary_keys','is_sendable','created','last_modified','tier','target_dmo'])
    for k,fl in sorted(byde.items(),key=lambda x:(layer(des[x[0]]['path']),des[x[0]]['name'])):
        d=des[k]; w.writerow([d['name'],k,d['path'],layer(d['path']),len(fl),';'.join(r['field'] for r in fl if r['pk']=='true'),'Y' if d['send']=='true' else '',d['created'],d['mod'],d['tier'],d['target']])

# Workbook
wb=Workbook(); hdr=Font(name='Arial',bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='1F3864'); body=Font(name='Arial')
def sheet(ws,headers,rows,widths):
    ws.append(headers)
    for c in ws[1]: c.font=hdr; c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='top')
    for row in rows: ws.append(row)
    for i,wd in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=wd
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=body; c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2'
cnt=Counter(d['tier'] for d in des.values())
ws=wb.active; ws.title='1_Strategy'
sheet(ws,['Principle','Detail'],[
 ['System of record first','Master data already ingested from CRM / source systems is not re-ingested from SFMC. Tier A = reconcile only.'],
 ['Ingest what SFMC owns','Engagement (already via MCE starter bundle), consent and subscription state, bounces, journey entry and send history, Einstein scores, marketing-acquired leads, Personalization behaviour.'],
 ['Standard DMO by default','Individual, ContactPointEmail/Phone/Address, Account, Lead, CommunicationSubscriptionConsent, EmailEngagement, MessageEngagement, Product / SalesOrder / ShoppingCart / WebsiteEngagement.'],
 ['Custom DMO only for four things','Journey_Entry, Journey_Send_Log, Einstein_Engagement_Score, Marketing_Exclusion.'],
 ['Consolidate before ingest','Union journey DEs with one SFMC SQL activity per custom DMO, then one data stream each. Fixes missing primary keys at the same time.'],
 ['Ingestion mechanism','MCE connector, Data Extension data streams (Connect API d360_datastream_create_sfmc_data_extension). Needs a PK and, for Engagement, a Date field.'],
 ['Identity','Map SFMC party key to the same ssot__Individual__dlm.ssot__Id__c the CRM connector uses. Retire any Demo IR ruleset first.'],
 ['Fix types upstream','Boolean flags typed Text, dates typed Text, emails and phones typed Text: change in SFMC before streaming.'],
 ['Reporting DEs','Do not ingest. Rebuild as Calculated Insights.'],
 ['Tier counts',', '.join(f'{k}: {v}' for k,v in sorted(cnt.items()))],
 ['Tier legend','A Reconcile only | B Bounce enrichment | C Consent and channel | D Leads | E Personalization | F Einstein | G Exclusions | H1 Journey entry | H2 Journey send log | I Rebuild as CI | X Skip']],[28,140])
ws=wb.create_sheet('2_DE_Mapping')
sheet(ws,['Tier','Approach','DE name','Customer key','Folder','Fields','Primary key','Target DMO','Rationale','Last modified'],
 [[d['tier'],d['approach'],d['name'],k,d['path'].replace('Data Extensions / ',''),len(byde[k]),';'.join(r['field'] for r in byde[k] if r['pk']=='true') or 'NONE',d['target'],d['rationale'],d['mod']] for k,d in sorted(des.items(),key=lambda x:(x[1]['tier'],x[1]['name']))],[6,16,42,38,40,8,28,46,70,12])
ws=wb.create_sheet('3_Field_Mapping'); rows=[]
for k,d in sorted(des.items(),key=lambda x:(x[1]['tier'],x[1]['name'])):
    if d['tier'] in ('A','I','X'): continue
    for r in sorted(byde[k],key=lambda x:int(x['ord'] or 0)):
        dmo,f,note=map_field(d,r); rows.append([d['tier'],d['name'],r['field'],r['ftype'],r['maxlen'],'Y' if r['pk']=='true' else '',dmo,f,note])
sheet(ws,['Tier','DE name','Source field','SFMC type','Max len','PK','Target DMO','Target field','Mapping note'],rows,[6,42,34,14,8,5,44,38,60])
ws=wb.create_sheet('4_Custom_DMO_Spec'); sheet(ws,['Custom DMO','Category','Primary key','Event / time field','Attributes','Source DEs'],rules['custom_dmos'],[30,12,44,20,90,60])
ws=wb.create_sheet('5_Gaps'); gaps=[]
for k,d in des.items():
    if d['tier'] in ('A','I','X'): continue
    fl=byde[k]
    if not any(r['pk']=='true' for r in fl): gaps.append(['No primary key',d['name'],'Wrap in a query DE with a composite key before creating the data stream'])
    for r in fl:
        n=r['field'].lower()
        if re.match(r'^(is|has)',n) and r['ftype']!='Boolean' and 'hashed' not in n: gaps.append(['Flag typed '+r['ftype'],d['name']+'.'+r['field'],'Change to Boolean in SFMC'])
        if re.search(r'(eventdate|sentdate|clickdate|birthday|deployment)',n) and r['ftype']=='Text': gaps.append(['Date typed Text',d['name']+'.'+r['field'],'Change to Date in SFMC'])
        if re.search(r'email',n) and r['ftype']=='Text' and not re.search(r'(opt|bounce|month|hashed|name|sender|deployment)',n): gaps.append(['Email typed Text',d['name']+'.'+r['field'],'Change to EmailAddress'])
alias=defaultdict(set)
for r in fields: alias[re.sub(r'[^a-z0-9]','',r['field'].lower())].add(r['field'])
for k,v in alias.items():
    if len(v)>1 and any(x in k for x in ('id','key','email')): gaps.append(['Key spelling variants',' / '.join(sorted(v)),'Alias table in loader so one mapping rule covers all'])
sheet(ws,['Gap type','Object / field','Fix'],gaps,[24,70,60])
wb.save(f'{out}/{prefix}_MCE_to_Data360_Mapping.xlsx')
print(json.dumps(dict(des=len(des),fields=len(fields),tiers=cnt,field_mapping_rows=len(rows),gaps=len(gaps)),indent=1))
